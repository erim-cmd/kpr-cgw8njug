"""Test syllabus'ları üretir (DOCX) — LibreOffice ile PDF'e çevrilir. Beklenen değerler expected.json'da."""
from docx import Document
from docx.shared import Pt
import json

def table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
    for i, r in enumerate(rows):
        for j, v in enumerate(r): t.cell(i, j).text = str(v)

exp = {}

# 1) Türkçe, tablolu BAU izlencesi
d = Document()
d.add_heading("BAHÇEŞEHİR ÜNİVERSİTESİ", 0)
d.add_paragraph("Mühendislik ve Doğa Bilimleri Fakültesi – Mekatronik Mühendisliği")
d.add_heading("DERS İZLENCESİ", 1)
table(d, [["Ders Kodu", "MCH 2016"], ["Ders Adı", "Devre Analizi"], ["Dönem", "2026-2027 Güz"],
          ["Kredi (T+U+K)", "3+0+3"], ["AKTS", "6"], ["Öğretim Üyesi", "Dr. Öğr. Üyesi Ayşe Yılmaz"],
          ["E-posta", "ayse.yilmaz@bau.edu.tr"], ["Ofis", "D-512"], ["Ofis Saatleri", "Çarşamba 14:00-16:00"],
          ["Ders Saatleri", "Pazartesi 09:00-10:50 (D-301), Perşembe 13:00-14:50 (Lab-2)"]])
d.add_heading("Değerlendirme", 2)
table(d, [["Etkinlik", "Sayı", "Katkı Payı (%)"], ["Ara Sınav (Vize)", "1", "30"], ["Ödevler", "4", "15"],
          ["Laboratuvar", "5", "15"], ["Final Sınavı", "1", "40"], ["Toplam", "", "100"]])
d.add_heading("Haftalık Plan", 2)
table(d, [["Hafta", "Tarih", "Konu", "Etkinlik"],
          ["1", "28.09.2026", "Temel kavramlar", ""],
          ["2", "05.10.2026", "Kirchhoff yasaları", ""],
          ["3", "12.10.2026", "Düğüm analizi", "Ödev 1 teslim"],
          ["5", "26.10.2026", "Thevenin ve Norton", "Ödev 2 teslim"],
          ["8", "16.11.2026", "Ara Sınav", "Vize Sınavı (18.11.2026, 10:00)"],
          ["10", "30.11.2026", "Kapasitör ve bobin", "Ödev 3 teslim"],
          ["12", "14.12.2026", "RLC devreleri", "Ödev 4 teslim"],
          ["15", "05.01.2027", "Final", "Final Sınavı: 12 Ocak 2027 13:00"]])
d.add_heading("Kurallar", 2)
d.add_paragraph("Derslerin en az %70'ine devam zorunludur. Devam şartını sağlamayan öğrenciler NA notu alır ve final sınavına giremez.")
d.add_paragraph("Final sınavından en az 40 puan alınması gerekmektedir; aksi halde öğrenci FF/F notu alır.")
d.add_paragraph("Geç teslim edilen ödevler kabul edilmez.")
d.add_paragraph("Mazeret sınavı yalnızca belgelendirilmiş sağlık raporu ile verilir.")
d.add_paragraph("Kopya veya intihal tespit edilmesi durumunda öğrenci disiplin kuruluna sevk edilir ve dersten F alır.")
d.save("tr_tablo.docx")
exp["tr_tablo"] = {"code": "MCH 2016", "name": "Devre Analizi", "instructor": "Dr. Öğr. Üyesi Ayşe Yılmaz",
  "email": "ayse.yilmaz@bau.edu.tr", "office": "D-512", "credit": 3, "ects": 6,
  "sessions": [[0, "09:00", "10:50", "D-301"], [3, "13:00", "14:50", "Lab-2"]],
  "grading": {"Ara Sınav (Vize)": 30, "Ödevler": 15, "Laboratuvar": 15, "Final Sınavı": 40},
  "attendance": 70, "final_min": 40,
  "items": [["Ödev 1", "2026-10-12"], ["Ödev 2", "2026-10-26"], ["Vize", "2026-11-18", "10:00"], ["Ödev 3", "2026-11-30"], ["Ödev 4", "2026-12-14"], ["Final", "2027-01-12", "13:00"]],
  "policies": ["devam", "baraj", "gec_teslim", "telafi", "durustluk"]}

# 2) İngilizce, paragraf + tarih yazıları
d = Document()
d.add_heading("MCH 3012 – Control Systems", 0)
d.add_paragraph("Bahcesehir University, Fall 2026")
d.add_paragraph("Instructor: Prof. Dr. Mehmet Kaya    Email: mehmet.kaya@eng.bau.edu.tr")
d.add_paragraph("Office: B-204    Office Hours: Tuesdays 11:00–12:00")
d.add_paragraph("Credits: 3    ECTS: 5")
d.add_paragraph("Lectures: Tuesday 13:00–15:50, Room A-105")
d.add_heading("Grading", 2)
d.add_paragraph("Quizzes (best 4 of 5): 10%")
d.add_paragraph("Homework: 20%")
d.add_paragraph("Midterm Exam: 30%")
d.add_paragraph("Final Exam: 40%")
d.add_heading("Important Dates", 2)
d.add_paragraph("Homework 1 due: October 13, 2026")
d.add_paragraph("Midterm Exam: November 17, 2026 at 13:00")
d.add_paragraph("Project Report due: December 22, 2026")
d.add_paragraph("Final Exam: TBA (during final exam week)")
d.add_heading("Policies", 2)
d.add_paragraph("Attendance: Students must attend at least 70% of the lectures; otherwise they will receive NA.")
d.add_paragraph("Late submissions will receive a 10% penalty per day.")
d.add_paragraph("The lowest quiz grade will be dropped.")
d.add_paragraph("Use of AI tools to write homework is considered plagiarism.")
d.save("en_paragraf.docx")
exp["en_paragraf"] = {"code": "MCH 3012", "name": "Control Systems", "instructor": "Prof. Dr. Mehmet Kaya",
  "email": "mehmet.kaya@eng.bau.edu.tr", "office": "B-204", "credit": 3, "ects": 5,
  "sessions": [[1, "13:00", "15:50", "A-105"]],
  "grading": {"Quizzes": 10, "Homework": 20, "Midterm Exam": 30, "Final Exam": 40},
  "attendance": 70, "final_min": None,
  "items": [["Homework 1", "2026-10-13"], ["Midterm", "2026-11-17", "13:00"], ["Project", "2026-12-22"], ["Final", ""]],
  "policies": ["devam", "gec_teslim", "not_kurali", "durustluk"]}

# 3) Türkçe, hafta numaralı, tarih yok, devamsızlık sayı ile
d = Document()
d.add_paragraph("Ders: EEE 2101 Sayısal Elektronik (Lojik Tasarım)")
d.add_paragraph("Yerel Kredi: 4    AKTS: 7")
d.add_paragraph("Öğretim Görevlisi: Öğr. Gör. Can Demir (can.demir@bau.edu.tr)")
d.add_paragraph("Ders günü ve saati: Salı 10:00-12:50, Cuma 14:00-15:50 / Derslik: C-201")
d.add_paragraph("Değerlendirme: Vize %35, Proje %25, Final %40")
d.add_paragraph("7. hafta: Vize sınavı")
d.add_paragraph("11. hafta: Proje sunumları")
d.add_paragraph("Final sınavı akademik takvimde ilan edilecektir.")
d.add_paragraph("Devamsızlık hakkı en fazla 4 derstir. Bütünleme sınavı final yerine geçer.")
d.save("tr_hafta.docx")
exp["tr_hafta"] = {"code": "EEE 2101", "name": "Sayısal Elektronik (Lojik Tasarım)", "instructor": "Öğr. Gör. Can Demir",
  "email": "can.demir@bau.edu.tr", "office": "", "credit": 4, "ects": 7,
  "sessions": [[1, "10:00", "12:50", "C-201"], [4, "14:00", "15:50", "C-201"]],
  "grading": {"Vize": 35, "Proje": 25, "Final": 40},
  "attendance": None, "max_absences": 4, "final_min": None,
  "items": [["Vize", "", None, 7], ["Proje", "", None, 11], ["Final", ""]],
  "policies": ["devam", "butunleme"]}

json.dump(exp, open("expected.json", "w"), ensure_ascii=False, indent=1)
